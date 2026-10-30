"""打码规则收紧：短名（≤6 字）保留前 2 字，长名保留前 6 字。

上一版对「上迈科技××」这种短名等于没打码 —— 名字本身就 4 个字，露出 4 个没意义。
"""
from __future__ import annotations

import collections
import json
import pathlib

from openpyxl import load_workbook

ROOT = pathlib.Path(r"C:\Users\31560\Documents\00_open\robowiki")
XLSX = pathlib.Path(r"C:\Users\31560\Downloads\【红黑榜】电子元件网店红黑榜.xlsx")
SUP = ROOT / "data/suppliers.json"

wb = load_workbook(XLSX, read_only=True, data_only=True)
ws = wb.worksheets[0]
counts: collections.Counter[str] = collections.Counter()
drop = set()
for r in list(ws.iter_rows(values_only=True))[2:]:
    if not r or not r[3]:
        continue
    if str(r[0] or "").strip() not in ("红榜", "黑榜"):
        continue
    counts[str(r[3]).strip()] += 1
wb.close()

# 已打码过的名字 → 原名（从 origin 里推不出来，改用「名字前缀匹配」恢复）
rows = json.loads(SUP.read_text(encoding="utf-8"))
names = list(counts)


def mask(name: str) -> str:
    keep = 2 if len(name) <= 6 else 6
    return name[:keep] + "××"


fixed = 0
for r in rows:
    if r.get("kind") != "component_shop" or not r.get("name_masked"):
        continue
    cur = r["name"]
    if not cur.endswith("××"):
        continue
    stem = cur[:-2]
    # 找原表里最匹配的原名
    cand = [n for n in names if n.startswith(stem)]
    if len(cand) != 1:
        continue
    real = cand[0]
    new = mask(real)
    if new != cur:
        r["name"] = new
        fixed += 1

SUP.write_text(json.dumps(rows, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
black = [r for r in rows if r.get("listing") == "黑榜"]
print(f"重新打码 {fixed} 家")
print("样例:", [r["name"] for r in black if r.get("name_masked")][:8])
print("保留全名:", [r["name"] for r in black if not r.get("name_masked")])
